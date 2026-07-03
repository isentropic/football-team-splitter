#!/usr/bin/env python3
import argparse
import datetime as dt
import json
import re
import shutil
import sqlite3
import uuid
from collections import Counter
from pathlib import Path

import openpyxl


SOURCE_NAMESPACE = uuid.UUID("2d4d8a9d-a62c-49d4-b638-bf1372c1de4f")
SOURCE_SHEET = "Games-2021-2025"
DEFAULT_XLSX = Path("/Users/yerlen/Downloads/Stats.xlsx")
DEFAULT_DB = Path(".wrangler/state/v3/d1/miniflare-D1DatabaseObject/87b878a1a36313b2ca548064a20cb25bdaf2a6038ff6eb9a2af98e62c8d37538.sqlite")
DEFAULT_CUTOFF = dt.date(2026, 1, 10)

COLORS = {"blue", "green", "orange", "white"}
STOP_WORDS = {
    "pts", "gd", "gf", "ga", "# of games", "w", "d", "l",
    "pts/game", "gd/game", "gf/game", "ga/game",
    "teams", "games", "score",
}
DATE_RE = re.compile(r"Game\s+(\d{1,2})[./](\d{1,2})[./](\d{2,4})")


def norm_name(name: str) -> str:
    return re.sub(r"\s+", " ", name.strip().lower().rstrip("."))


def clean_player_name(name: str) -> str:
    return re.sub(r"\s+\([cC]\)$", "", name.strip())


def stable_id(*parts: object) -> str:
    return str(uuid.uuid5(SOURCE_NAMESPACE, ":".join(str(part) for part in parts)))


def parse_date(value) -> dt.date | None:
    match = DATE_RE.search(str(value or ""))
    if not match:
        return None
    day, month, year = [int(part) for part in match.groups()]
    if year < 100:
        year += 2000
    return dt.date(year, month, day)


def local_noon_ms(day: dt.date, session_index_for_day: int) -> int:
    # Match the app's millisecond timestamp style while keeping repeated same-day
    # historical blocks in a stable order.
    when = dt.datetime(day.year, day.month, day.day, 12, 0) + dt.timedelta(minutes=30 * session_index_for_day)
    return int(when.timestamp() * 1000)


def read_cell(ws, row: int, col: int):
    return ws.cell(row, col).value


def parse_block(ws, start_row: int) -> dict:
    session_date = parse_date(read_cell(ws, start_row, 2))
    if session_date is None:
        raise ValueError(f"No date at row {start_row}")

    score_cols: list[tuple[str, int]] = []
    for col in range(4, 12):
        value = read_cell(ws, start_row + 1, col)
        if isinstance(value, str) and value.strip().lower() in COLORS:
            score_cols.append((value.strip().lower(), col))
        elif score_cols:
            break

    roster_header_row = None
    team_cols: list[tuple[str, int]] = []
    for row in range(start_row + 1, start_row + 5):
        found: list[tuple[str, int]] = []
        for col in range(1, ws.max_column + 1):
            value = read_cell(ws, row, col)
            if isinstance(value, str) and value.strip().lower() in COLORS:
                found.append((value.strip().lower(), col))
        right_side = [item for item in found if item[1] > 6]
        if len(right_side) >= len(score_cols):
            roster_header_row = row
            team_cols = right_side[:len(score_cols)]
            break

    if not score_cols or roster_header_row is None or not team_cols:
        raise ValueError(f"Could not locate score/team columns at row {start_row}")

    teams = []
    for color, col in team_cols:
        players = []
        for row in range(roster_header_row + 1, min(start_row + 20, ws.max_row + 1)):
            value = read_cell(ws, row, col)
            if value is None:
                continue
            if not isinstance(value, str):
                continue
            text = value.strip()
            lower = text.lower()
            if not text:
                continue
            if lower in STOP_WORDS or lower in COLORS:
                break
            players.append(clean_player_name(text))
        teams.append({"color": color, "players": players})

    games = []
    for row in range(start_row + 2, start_row + 20):
        game_no = read_cell(ws, row, 3)
        if not isinstance(game_no, (int, float)):
            continue
        scores = []
        for color, col in score_cols:
            value = read_cell(ws, row, col)
            if isinstance(value, (int, float)):
                scores.append((color, int(value)))
        if len(scores) == 2:
            games.append({
                "game_no": int(game_no),
                "team1": scores[0][0],
                "score1": scores[0][1],
                "team2": scores[1][0],
                "score2": scores[1][1],
            })

    return {
        "source_sheet": ws.title,
        "source_row": start_row,
        "date": session_date,
        "teams": teams,
        "games": games,
    }


def parse_workbook(path: Path, sheet_name: str, cutoff: dt.date) -> list[dict]:
    wb = openpyxl.load_workbook(path, data_only=True, read_only=False)
    ws = wb[sheet_name]
    sessions = []
    for row in range(1, ws.max_row + 1):
        session_date = parse_date(read_cell(ws, row, 2))
        if session_date and session_date < cutoff:
            parsed = parse_block(ws, row)
            if parsed["games"]:
                sessions.append(parsed)
    return sessions


def load_players(conn: sqlite3.Connection) -> tuple[dict[str, str], set[str]]:
    rows = conn.execute("SELECT id, name FROM players").fetchall()
    return {norm_name(name): player_id for player_id, name in rows}, {name for _, name in rows}


def make_import_rows(conn: sqlite3.Connection, sessions: list[dict]) -> dict:
    existing_name_to_id, existing_names = load_players(conn)
    all_names = sorted({
        name
        for session in sessions
        for team in session["teams"]
        for name in team["players"]
    }, key=str.lower)

    name_to_id: dict[str, str] = {}
    created_players = []
    matched_players = []
    for name in all_names:
        normalized = norm_name(name)
        existing_id = existing_name_to_id.get(normalized)
        if existing_id:
            name_to_id[name] = existing_id
            matched_players.append(name)
        else:
            player_id = stable_id("player", normalized)
            name_to_id[name] = player_id
            created_players.append((player_id, name))

    by_date_counter: Counter[dt.date] = Counter()
    session_rows = []
    game_rows = []
    for session in sorted(sessions, key=lambda item: (item["date"], item["source_row"])):
        session_index_for_day = by_date_counter[session["date"]]
        by_date_counter[session["date"]] += 1
        played_at = local_noon_ms(session["date"], session_index_for_day)
        session_id = stable_id(SOURCE_SHEET, session["source_row"], session["date"].isoformat())
        teams = [
            {
                "color": team["color"],
                "playerIds": [name_to_id[name] for name in team["players"]],
            }
            for team in session["teams"]
        ]
        session_rows.append((session_id, json.dumps(teams, separators=(",", ":")), played_at, session))

        for game in session["games"]:
            game_id = stable_id(SOURCE_SHEET, session["source_row"], session["date"].isoformat(), "game", game["game_no"])
            game_played_at = played_at + game["game_no"] * 60_000
            game_rows.append((
                game_id,
                session_id,
                game["team1"],
                game["score1"],
                game["team2"],
                game["score2"],
                game_played_at,
                session,
                game,
            ))

    return {
        "matched_players": matched_players,
        "created_players": created_players,
        "historical_player_ids": [stable_id("player", norm_name(name)) for name in all_names],
        "session_rows": session_rows,
        "game_rows": game_rows,
        "existing_player_names": existing_names,
    }


def print_report(conn: sqlite3.Connection, sessions: list[dict], rows: dict) -> None:
    existing_sessions = conn.execute("SELECT COUNT(*) FROM sessions").fetchone()[0]
    existing_games = conn.execute("SELECT COUNT(*) FROM games").fetchone()[0]
    existing_players = conn.execute("SELECT COUNT(*) FROM players").fetchone()[0]
    print("Import preview")
    print(f"  source sheet: {SOURCE_SHEET}")
    print(f"  sessions parsed before {DEFAULT_CUTOFF.isoformat()}: {len(sessions)}")
    print(f"  games parsed: {len(rows['game_rows'])}")
    print(f"  players currently in local DB: {existing_players}")
    print(f"  existing local sessions/games: {existing_sessions}/{existing_games}")
    print(f"  matched historical player names: {len(rows['matched_players'])}")
    print(f"  new historical players to create: {len(rows['created_players'])}")
    if rows["created_players"]:
        print("  first new players:")
        for _, name in rows["created_players"][:40]:
            print(f"    - {name}")
    print("  first sessions:")
    for session_id, _, played_at, source in rows["session_rows"][:5]:
        day = dt.datetime.fromtimestamp(played_at / 1000).strftime("%Y-%m-%d %H:%M")
        print(f"    - {day} row {source['source_row']} id {session_id[:8]} games {len(source['games'])}")
    print("  last sessions:")
    for session_id, _, played_at, source in rows["session_rows"][-5:]:
        day = dt.datetime.fromtimestamp(played_at / 1000).strftime("%Y-%m-%d %H:%M")
        print(f"    - {day} row {source['source_row']} id {session_id[:8]} games {len(source['games'])}")


def apply_import(conn: sqlite3.Connection, db_path: Path, rows: dict) -> None:
    backup_dir = Path(".wrangler/backups")
    backup_dir.mkdir(parents=True, exist_ok=True)
    backup_path = backup_dir / f"local-d1-before-xlsx-history-{dt.datetime.now().strftime('%Y%m%d-%H%M%S')}.sqlite"
    shutil.copy2(db_path, backup_path)
    print(f"Backup created: {backup_path}")

    with conn:
        conn.executemany(
            """
            INSERT OR IGNORE INTO players
              (id, name, pace, shooting, passing, dribbling, defending, physique, morale, retired)
            VALUES (?, ?, 5, 5, 5, 5, 5, 5, 5, 1)
            """,
            rows["created_players"],
        )
        conn.executemany(
            "UPDATE players SET retired = 1 WHERE id = ?",
            [(player_id,) for player_id in rows["historical_player_ids"]],
        )
        conn.executemany(
            "INSERT OR IGNORE INTO sessions (id, teams, played_at) VALUES (?, ?, ?)",
            [(sid, teams, played_at) for sid, teams, played_at, _ in rows["session_rows"]],
        )
        conn.executemany(
            """
            INSERT OR IGNORE INTO games
              (id, session_id, team1, score1, team2, score2, played_at)
            VALUES (?, ?, ?, ?, ?, ?, ?)
            """,
            [(gid, sid, t1, s1, t2, s2, played_at) for gid, sid, t1, s1, t2, s2, played_at, _, _ in rows["game_rows"]],
        )


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--xlsx", type=Path, default=DEFAULT_XLSX)
    parser.add_argument("--db", type=Path, default=DEFAULT_DB)
    parser.add_argument("--apply", action="store_true")
    args = parser.parse_args()

    sessions = parse_workbook(args.xlsx, SOURCE_SHEET, DEFAULT_CUTOFF)
    conn = sqlite3.connect(args.db)
    try:
        rows = make_import_rows(conn, sessions)
        print_report(conn, sessions, rows)
        if args.apply:
            before = {
                table: conn.execute(f"SELECT COUNT(*) FROM {table}").fetchone()[0]
                for table in ("players", "sessions", "games")
            }
            apply_import(conn, args.db, rows)
            after = {
                table: conn.execute(f"SELECT COUNT(*) FROM {table}").fetchone()[0]
                for table in ("players", "sessions", "games")
            }
            print("Applied local import")
            for table in ("players", "sessions", "games"):
                print(f"  {table}: {before[table]} -> {after[table]} (+{after[table] - before[table]})")
        else:
            print("Dry run only. Use --apply to write to local D1 sqlite.")
    finally:
        conn.close()


if __name__ == "__main__":
    main()
